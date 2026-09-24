"""Excel export in the IDN onboarding requirements template format."""

from __future__ import annotations

import io
from datetime import datetime, timezone
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

import interview

GREEN = PatternFill("solid", fgColor="92D050")
CYAN = PatternFill("solid", fgColor="00B0F0")
TAN = PatternFill("solid", fgColor="DDD9C4")
GREY = PatternFill("solid", fgColor="D9D9D9")
BOLD = Font(bold=True, size=10)
NORMAL = Font(size=10)
ITALIC = Font(italic=True, size=9)
THIN = Side(style="thin", color="000000")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
WRAP_RIGHT = Alignment(wrap_text=True, vertical="center", horizontal="right")
CENTRE = Alignment(horizontal="center", vertical="center", wrap_text=True)


def answer(state: dict[str, Any], *slot_ids: str) -> str:
    """First non-empty answer among the given slots, or a follow-up note."""
    for slot_id in slot_ids:
        entry = state["answers"].get(slot_id)
        if entry:
            return interview.render(entry["value"])
        if slot_id in state["open_questions"]:
            reason = state["open_questions"][slot_id].get("reason") or "unassigned"
            return f"[To be confirmed - {reason}]"
    return ""


def table(state: dict[str, Any], slot_id: str) -> list[dict[str, Any]]:
    entry = state["answers"].get(slot_id)
    value = entry["value"] if entry else None
    return value if isinstance(value, list) and value and isinstance(value[0], dict) else []


def band(ws: Worksheet, row: int, text: str, span: int, fill: PatternFill = GREEN) -> int:
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=span)
    cell = ws.cell(row=row, column=1, value=text)
    cell.fill = fill
    cell.font = BOLD
    cell.alignment = CENTRE
    for col in range(1, span + 1):
        ws.cell(row=row, column=col).border = BORDER
    return row + 1


def headers(ws: Worksheet, row: int, labels: list[str]) -> int:
    for col, label in enumerate(labels, start=1):
        cell = ws.cell(row=row, column=col, value=label)
        cell.fill = CYAN
        cell.font = BOLD
        cell.alignment = CENTRE
        cell.border = BORDER
    return row + 1


def data_row(ws: Worksheet, row: int, values: list[Any], fill: PatternFill | None = None) -> int:
    for col, value in enumerate(values, start=1):
        cell = ws.cell(row=row, column=col, value=value)
        cell.font = NORMAL
        cell.border = BORDER
        cell.alignment = WRAP_RIGHT if col == 2 else WRAP
        if fill:
            cell.fill = fill
    return row + 1


def widths(ws: Worksheet, values: list[int]) -> None:
    for index, width in enumerate(values, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


# --------------------------------------------------------------- sheet 1

# (item text, slot ids feeding the Requirement column)
PART_1 = [
    ("Is Birthright Provisioning required?\n(Access granted to all new hires)", ("birthright_required", "joiner_rules")),
    (
        "What account attributes (Firstname, Lastname, etc.) and other application "
        "information (role, access type, region) is needed to create an account?",
        ("user_record_fields",),
    ),
]
PART_2 = [
    ("How are movers handled today?", ("current_admin_process",)),
    ("How is access deprovisioned during mover event? Ex: Company Change", ("mover_rules",)),
    ("How is access deprovisioned during mover event? Ex: Title + Manager Change", ("mover_rules",)),
]
PART_3 = [
    ("How is access termination requested today?", ("current_admin_process",)),
    (
        "How is access deprovisioned during leaver events? What actions are taken "
        "during the leaver process?",
        ("leaver_rules", "retention_period"),
    ),
]
PART_4 = [
    (
        "When is the User Access Review cycle for this application? "
        "(Monthly, Quarterly, Semi-Annual, Annual, etc.)",
        ("certification_cadence",),
    ),
    (
        "Is the review performed on all the access a user has in the application, "
        "only on privileged access?",
        ("privileged_levels",),
    ),
    ("Who reviews the user's access in the application?", ("access_level_owner",)),
]
PART_5 = [
    ("Entitlements to be requested", ("access_levels_list",)),
    ("Roles to be requested", ("roles_to_request",)),
    ("Access Models/Profiles to be requested", ("access_profiles_to_request",)),
    ("Request only for existing users? No new user account creation?", ("existing_users_only",)),
]
PART_6 = [
    ("Manager Approval", ("approval_chain",)),
    ("Business Owner Approval", ("business_owner",)),
    ("Third Level Approval", ("segregation_of_duties",)),
]
PART_7 = [
    ("ServiceNow Fulfillment Group", ("snow_fulfillment_group",)),
    ("ServiceNow Assignment Group", ("snow_assignment_group",)),
    ("ServiceNow Assignment Group Manager", ("snow_assignment_group_manager",)),
    ("ServiceNow Service", ("snow_service",)),
    ("ServiceNow Service Offering Sys_ID", ("snow_service_offering_sysid",)),
    ("Automated Provisioning", ("automation_wanted", "integration_methods")),
]
PART_8 = [
    ("Comments", ("business_purpose",)),
    ("Static Applications", ("static_application",)),
    ("Mover", ("mover_rules",)),
    ("SNOW Ticket (if Required)", ("snow_ticket_required",)),
    ("Testing", ("test_identities",)),
]


def sheet_requirements(wb: Workbook, state: dict[str, Any]) -> None:
    ws = wb.create_sheet("IDN Requirements")
    widths(ws, [6, 46, 44, 24, 32])
    row = 1

    row = band(ws, row, "Introduction", 5)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
    intro = ws.cell(
        row=row,
        column=1,
        value=(
            f"This tab details the requirements necessary to on-board "
            f"{state['name']} into SailPoint Identity Security Cloud for the purpose of "
            "supporting identity onboarding/offboarding, access requests and access reviews."
        ),
    )
    intro.font = NORMAL
    intro.alignment = WRAP
    ws.row_dimensions[row].height = 28
    row += 2

    number = 1
    four_col = ["#", "Item", "Requirement", "Comments", ""]
    blocks = [
        ("Part - 1 [Joiner]", PART_1, four_col, None),
        ("Part - 2 [Mover]", PART_2, four_col, None),
        ("Part - 3 [Leaver]", PART_3, four_col, None),
        ("Part - 4 [Access Review]", PART_4, four_col, None),
        ("Part - 5 [Access Request]", PART_5, four_col, None),
        ("Part - 6 [Approval Process]", PART_6, ["#", "Item", "Requirement", "Name", "Comments"], None),
        ("Part - 7 [Fulfillment Details]", PART_7, ["#", "Item", "Description", "Name", "Comments"], TAN),
        ("Part - 8 [Assumptions]", PART_8, ["#", "Item", "Description", "Responsibility", "Comments"], None),
    ]

    for title, items, cols, fill in blocks:
        row = band(ws, row, title, 5)
        row = headers(ws, row, cols)
        for item, slots in items:
            row = data_row(ws, row, [number, item, answer(state, *slots), "", ""], fill)
            number += 1
        row += 1

    row = band(ws, row, "RFA Considerations", 5)
    row = headers(ws, row, ["Question", "Comments", "", "", ""])
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
    row = data_row(
        ws,
        row,
        [
            "How is the form defined currently in RFA and how it should look in ISC?",
            answer(state, "rfa_form_definition"),
        ],
    )
    ws.row_dimensions[row - 1].height = 46
    row += 1

    row = band(ws, row, "Acceptance Criteria", 5)
    row = headers(ws, row, ["Item", "Requirement Number", "", "", ""])
    criteria = answer(state, "acceptance_criteria") or "\n".join(
        f"{i}) {text}"
        for i, text in enumerate(
            filter(
                None,
                [
                    answer(state, "joiner_rules"),
                    answer(state, "mover_rules"),
                    answer(state, "leaver_rules"),
                    answer(state, "approval_chain"),
                    answer(state, "certification_cadence"),
                ],
            ),
            start=1,
        )
    )
    ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
    row = data_row(ws, row, ["ISC Requirements", criteria])
    ws.row_dimensions[row - 1].height = 90


# --------------------------------------------------------------- sheet 2

# Operation name -> the api_operations answer that makes it Required
OPERATION_TRIGGERS = {
    "Test Connection": None,  # always required
    "Account Aggregation": "List all users",
    "Account Delta Aggregation": None,
    "Group Aggregation": "List roles / groups",
    "Get Object": "Get a single user",
    "Get Object-Group": None,
    "Create Account": "Create a user",
    "Update Account": "Update a user",
    "Delete Account": "Delete a user",
    "Enable Account": "Disable / deactivate a user",
    "Disable Account": "Disable / deactivate a user",
    "Unlock Account": None,
    "Change Password": None,
    "Add Entitlement": "Assign a role / group",
    "Remove Entitlement": "Remove a role / group",
    "Partitioned Account Aggregation": None,
}

UAR_FIELDS = [
    ("Employee ID [EMP_ID/TSCID]", "G99999", "employeeID/TSCID from the HR system", "yes"),
    ("Account ID", "", "if different from correlation attribute", ""),
    ("First Name", "", "", "Good to have"),
    ("Last Name", "", "", "Good to have"),
    ("Status", "Active/Inactive", "status of account", "yes"),
    (
        "Additional fields can be added based on app/data requirements",
        "",
        "Other attributes required to make a decision during access review process",
        "",
    ),
    ("Entitlement Type", "Either comma separated values or values in separate rows", "", "yes"),
]


def sheet_web_services(wb: Workbook, state: dict[str, Any]) -> None:
    ws = wb.create_sheet("Web Services")
    widths(ws, [30, 26, 26, 24, 22, 18, 22, 22, 20])
    row = 1

    row = band(ws, row, "Instructions", 9)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
    note = ws.cell(
        row=row,
        column=1,
        value=(
            "As part of application onboarding, provide the information below. "
            "The Web Services connector is the preferred integration pattern for this application."
        ),
    )
    note.font = NORMAL
    note.alignment = WRAP
    row += 2

    row = band(ws, row, "Approximate number of users for the application", 4)
    row = data_row(ws, row, ["Approximate Number of Users:", answer(state, "user_volume")])
    row += 2

    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
    prompt = ws.cell(row=row, column=1, value="Please fill out the following data for accessing the web services per environment:")
    prompt.font = NORMAL
    row += 1

    row = headers(
        ws,
        row,
        [
            "Environment",
            "Server DNS Name",
            "Server IP Address",
            "Authentication Type\n(Basic/Oauth/API Token)",
            "Base URL",
            "Grant Type",
            "Token URL",
            "Username/Client ID",
            "Custom Config",
        ],
    )
    auth = answer(state, "auth_method")
    grid = {r.get("Environment", ""): r for r in table(state, "ws_environment_matrix")}
    env_note = answer(state, "environment_list")
    for env in ["Dev", "Test", "UAT", "PROD"]:
        r = grid.get(env, {})
        row = data_row(
            ws,
            row,
            [
                env,
                r.get("Server DNS Name", ""),
                r.get("Server IP Address", ""),
                r.get("Authentication Type") or (auth if env == "PROD" else ""),
                r.get("Base URL", ""),
                r.get("Grant Type") or (answer(state, "oauth_grant_type") if env == "PROD" else ""),
                r.get("Token URL") or (answer(state, "oauth_token_url") if env == "PROD" else ""),
                r.get("Username/Client ID") or (answer(state, "api_client_id") if env == "PROD" else ""),
                r.get("Custom Config", ""),
            ],
        )
    if env_note:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=9)
        cell = ws.cell(row=row, column=1, value=f"Reported environments: {env_note}")
        cell.font = ITALIC
        cell.alignment = WRAP
        row += 1
    row += 1

    row = band(ws, row, "Context URLs for different Operation", 5, CYAN)
    row = data_row(ws, row, ["API Documentation URL", answer(state, "api_documentation")])
    row += 1

    row = headers(ws, row, ["Operation", "Context URL", "Sample Request", "Sample Response", "Required?"])
    selected = state["answers"].get("api_operations", {}).get("value", []) or []
    ops = {r.get("Operation", ""): r for r in table(state, "ws_operation_matrix")}
    endpoints = answer(state, "ws_endpoint_inventory")
    for operation, trigger in OPERATION_TRIGGERS.items():
        entered = ops.get(operation, {})
        stated = str(entered.get("Required?", "")).strip()
        if stated:
            required = stated
        elif operation in ("Test Connection", "Account Aggregation"):
            required = "Required"
        elif operation == "Account Delta Aggregation":
            required = "Required" if answer(state, "ws_delta_support").startswith("Yes") else "Not Required"
        elif trigger:
            required = "Required" if trigger in selected else "Not Required"
        else:
            required = "Not Required"
        row = data_row(
            ws,
            row,
            [
                operation,
                entered.get("Context URL", ""),
                entered.get("Sample Request", ""),
                entered.get("Sample Response", ""),
                required,
            ],
        )
    if endpoints:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        cell = ws.cell(row=row, column=1, value=f"Endpoints reported at intake: {endpoints}")
        cell.font = ITALIC
        cell.alignment = WRAP
        ws.row_dimensions[row].height = 40
        row += 1
    row += 1

    row = band(ws, row, "User Access Review Data Query", 4)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
    cell = ws.cell(
        row=row,
        column=1,
        value=(
            "This API provides users and their associated entitlements to be used for "
            "reviewing access. At least one of the following three unique identifiers will be required."
        ),
    )
    cell.font = NORMAL
    cell.alignment = WRAP
    row += 2

    row = headers(ws, row, ["Field", "Example", "Definition", "Required?", ""])
    grid = {r.get("Field", ""): r for r in table(state, "uar_field_matrix")}
    identifier = answer(state, "corporate_id_field", "unique_identifier")
    account_id = answer(state, "account_id_field")
    status = answer(state, "inactive_representation")
    for field, example, definition, required in UAR_FIELDS:
        key = field.split(" [")[0]
        entered = grid.get(key, {})
        if entered.get("Field name in the application"):
            definition = entered["Field name in the application"]
        elif field.startswith("Employee ID") and identifier:
            definition = identifier
        elif field == "Account ID" and account_id:
            definition = account_id
        elif field == "Status" and status:
            definition = status
        row = data_row(
            ws,
            row,
            [field, entered.get("Example value") or example, definition, required, entered.get("Notes", "")],
        )
    extra = answer(state, "user_record_fields")
    if extra:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        cell = ws.cell(row=row, column=1, value=f"Fields reported at intake: {extra}")
        cell.font = ITALIC
        cell.alignment = WRAP
        ws.row_dimensions[row].height = 32
        row += 1
    row += 1

    row = band(ws, row, "Entitlement Information (separate query for each entitlement with proper description)", 3)
    row = headers(ws, row, ["Entitlement Type", "Entitlement Name", "Entitlement Description", "", ""])
    rows = table(state, "entitlement_matrix")
    if rows:
        for r in rows:
            row = data_row(
                ws,
                row,
                [
                    r.get("Entitlement Type", ""),
                    r.get("Entitlement Name", ""),
                    r.get("Entitlement Description", ""),
                    "",
                    "",
                ],
            )
    else:
        levels = answer(state, "access_levels_list")
        lines = [line.strip() for line in levels.splitlines() if line.strip()] if levels else []
        for line in lines or ["", ""]:
            name, _, description = line.partition("-")
            row = data_row(ws, row, ["Group", name.strip(), description.strip(), "", ""])
    for _ in range(2):
        row = data_row(ws, row, ["", "", "", "", ""])


# --------------------------------------------------------------- sheet 3

def sheet_intake(wb: Workbook, schema: dict[str, Any], state: dict[str, Any]) -> None:
    """Raw captured answers, so every cell above is traceable to a question."""
    ws = wb.create_sheet("Intake record")
    widths(ws, [24, 52, 54, 34])
    sections = {s["id"]: s["name"] for s in schema["sections"]}
    by_id = {s["id"]: s for s in schema["slots"]}
    order = {s["id"]: i for i, s in enumerate(schema["sections"])}

    row = headers(ws, 1, ["Section", "Question", "Answer", "Flagged for review"])
    rows = sorted(
        state["answers"].items(), key=lambda kv: order.get(by_id[kv[0]]["section"], 99)
    )
    for slot_id, entry in rows:
        slot = by_id[slot_id]
        row = data_row(
            ws,
            row,
            [
                sections[slot["section"]],
                slot["question"],
                interview.render(entry["value"]),
                entry.get("review", ""),
            ],
        )
    ws.freeze_panes = "A2"

    if state["open_questions"]:
        ws2 = wb.create_sheet("Open questions")
        widths(ws2, [24, 60, 38])
        r = headers(ws2, 1, ["Section", "Question", "Who will follow up"])
        for slot_id, item in state["open_questions"].items():
            slot = by_id[slot_id]
            r = data_row(
                ws2,
                r,
                [sections[slot["section"]], slot["question"], item.get("reason") or "unassigned"],
            )


def build(schema: dict[str, Any], state: dict[str, Any]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)

    sheet_requirements(wb, state)

    methods = state["answers"].get("integration_methods", {}).get("value", []) or []
    if any(m in methods for m in ("SCIM 2.0 API", "REST / JSON API", "SOAP / XML API")):
        sheet_web_services(wb, state)

    sheet_intake(wb, schema, state)

    summary = wb.create_sheet("Document info", 0)
    widths(summary, [28, 60])
    r = headers(summary, 1, ["Field", "Value"])
    for label, value in [
        ("Application", state["name"]),
        ("Reference", state["slug"]),
        ("Intake started", state.get("created", "")),
        ("Submitted", state.get("submitted_at", "not submitted")),
        ("Exported", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
        ("Answered", str(len(state["answers"]))),
        ("Flagged for follow-up", str(len(state["open_questions"]))),
    ]:
        r = data_row(summary, r, [label, value])

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
